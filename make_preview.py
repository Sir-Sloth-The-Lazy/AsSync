"""
make_preview.py — Build an Excel preview of exactly what will land in Asana.

Useful before the first real sync: it lets the team eyeball 400+ rows in a
familiar tool instead of reading JSON.

    python make_preview.py --input inbox --out build/asana_preview.xlsx
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from parse_rise import parse_folder

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=10)
RAG = {"Green": "C6EFCE", "Yellow": "FFEB9C", "Red": "FFC7CE"}
SEV = {"error": "FFC7CE", "warning": "FFEB9C"}


def write_header(ws, headers):
    """Must run before any append() — on an empty sheet append() writes row 1."""
    ws.append(headers)


def style_sheet(ws, headers, widths, n_rows, table_name):
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    if n_rows:
        ref = f"A1:{get_column_letter(len(headers))}{n_rows + 1}"
        t = Table(displayName=table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", type=Path, default=Path("inbox"))
    ap.add_argument("--out", type=Path, default=Path("build/asana_preview.xlsx"))
    args = ap.parse_args()

    records, issues = parse_folder(args.input)
    wb = Workbook()

    # ---- Milestones ----
    ws = wb.active
    ws.title = "Milestone tasks"
    hdr = ["Function", "Priority", "Section", "Task name", "Start", "Due",
           "Week", "KPI", "Owner in sheet", "RAG", "Source cell", "RISE Source ID"]
    write_header(ws, hdr)
    ms = [r for r in records if r.record_kind == "milestone"]
    ms.sort(key=lambda r: (r.function, r.priority_no, r.week_no or 0))
    for r in ms:
        ws.append([r.function, f"P{r.priority_no}", f"Priority {r.priority_no}",
                   r.title, r.start_on, r.week_end, r.week_no,
                   (r.kpi_name or "")[:90], r.kpi_owner or r.priority_owner or "",
                   r.kpi_status or r.priority_status or "", r.source_cell, r.record_id])
    style_sheet(ws, hdr, [16, 9, 13, 62, 11, 11, 7, 44, 18, 9, 11, 19], len(ms), "Milestones")
    for row in ws.iter_rows(min_row=2, max_row=len(ms) + 1):
        for c in row:
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (4, 8))
        rag = row[9].value
        if rag in RAG:
            row[9].fill = PatternFill("solid", fgColor=RAG[rag])

    # ---- KPIs ----
    ws2 = wb.create_sheet("KPI tasks")
    hdr2 = ["Function", "Priority", "Task name", "Type", "Goal", "RAG",
            "Readings", "Green", "Yellow", "Red", "RISE Source ID"]
    write_header(ws2, hdr2)
    kp = [r for r in records if r.record_kind == "kpi"]
    kp.sort(key=lambda r: (r.function, r.priority_no))
    for r in kp:
        reads = "  ".join(f"{k}:{v}" for k, v in
                          sorted(r.measurements.items(), key=lambda kv: int(kv[0][1:])))
        th = r.thresholds or {}
        ws2.append([r.function, f"P{r.priority_no}", f"KPI · {r.title}", r.kpi_type or "",
                    (r.goal or "")[:70], r.kpi_status or "", reads[:150],
                    (th.get("green") or "")[:40], (th.get("yellow") or "")[:40],
                    (th.get("red") or "")[:40], r.record_id])
    style_sheet(ws2, hdr2, [16, 9, 58, 10, 40, 9, 46, 24, 24, 24, 19], len(kp), "KPIs")
    for row in ws2.iter_rows(min_row=2, max_row=len(kp) + 1):
        for c in row:
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 5, 7))
        if row[5].value in RAG:
            row[5].fill = PatternFill("solid", fgColor=RAG[row[5].value])

    # ---- Data quality ----
    ws3 = wb.create_sheet("Data quality")
    hdr3 = ["Severity", "File", "Sheet", "Finding"]
    write_header(ws3, hdr3)
    order = {"error": 0, "warning": 1}
    for i in sorted(issues, key=lambda x: (order[x.severity], x.source_file)):
        ws3.append([i.severity.upper(), i.source_file, i.source_sheet or "", i.message])
    style_sheet(ws3, hdr3, [11, 38, 14, 96], len(issues), "Issues")
    for row in ws3.iter_rows(min_row=2, max_row=len(issues) + 1):
        for c in row:
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=c.column == 4)
        row[0].fill = PatternFill("solid", fgColor=SEV[row[0].value.lower()])

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out}  ({len(ms)} milestones, {len(kp)} KPIs, {len(issues)} findings)")


if __name__ == "__main__":
    main()
